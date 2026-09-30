import bz2
import gzip
import io
import json
import lzma
import shutil
import stat
import subprocess
import tarfile
import wave
import zipfile

import pytest
from openpyxl import Workbook

from evidencekg.config import DEFAULTS, initialize
from evidencekg.ingest import ingest
from evidencekg.parsers import parse, signature
from evidencekg.retrieval import API


def run(data, suffix="", **options):
    return parse(data, suffix, {**DEFAULTS, "ocr": "off", **options})


def text(result):
    return "\n".join(s["text"] for s in result["sections"])


def zipped(entries):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries:
            archive.writestr(name, data)
    return output.getvalue()


@pytest.mark.parametrize("suffix", ["", ".weird", ".csv", ".json", ".py", ".yaml"])
def test_text_detection(suffix):
    result = run(b"Fictional searchable text", suffix)
    assert result["status"] == "ready"
    assert "searchable" in text(result)


def test_binary_is_not_text():
    assert run(b"\x00\x01binary", "")["status"] == "unsupported"


def test_html_calendar_csv_and_rtf():
    assert "Hello" in text(run(b"<html><script>secret()</script><p>Hello</p></html>"))
    assert "secret" not in text(run(b"<html><script>secret()</script><p>Hello</p></html>"))
    assert "meeting today" in text(run(b"BEGIN:VCALENDAR\r\nSUMMARY:meeting\r\n  today\r\nEND:VCALENDAR"))
    result = run(b'name,description\nAlice,"two\nlines"', ".csv")
    assert len(result["sections"]) == 2
    assert json.loads(result["sections"][1]["text"]) == ["Alice", "two\nlines"]
    assert "Hello" in text(run(b"{\\rtf1\\ansi Hello \\b world\\b0}", ".rtf"))


@pytest.mark.parametrize("suffix", ["", ".xlsx", ".xlst", ".xltx", ".xlsm", ".xltm"])
def test_spreadsheet_cells_and_formulas(suffix):
    book = Workbook()
    book.active.title = "Evidence"
    book.active["B2"] = "Fictional invoice"
    book.active["C2"] = "=1+1"
    output = io.BytesIO()
    book.save(output)
    result = run(output.getvalue(), suffix)
    assert result["status"] == "partial"
    assert "Fictional invoice" in text(result)
    assert "=1+1" in text(result)
    assert any(s["locator"].get("cell") == "B2" for s in result["sections"])


@pytest.mark.parametrize("suffix", ["", ".docx", ".dotx", ".doct", ".docm", ".dotm"])
def test_word_and_templates(suffix):
    from docx import Document

    doc = Document()
    doc.add_paragraph("Fictional letter")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table contents"
    output = io.BytesIO()
    doc.save(output)
    data = output.getvalue()
    if suffix == ".dotx":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            data = zipped(
                [
                    (
                        n,
                        archive.read(n).replace(
                            b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                            b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
                        )
                        if n == "[Content_Types].xml"
                        else archive.read(n),
                    )
                    for n in archive.namelist()
                ]
            )
    result = run(data, suffix)
    assert result["status"] == "partial"
    assert "Fictional letter" in text(result)
    assert "Table contents" in text(result)


@pytest.mark.parametrize(
    "suffix,entries,expected",
    [
        (".odt", [("content.xml", "<root><p>Open document</p></root>")], "Open document"),
        (
            ".pptx",
            [
                ("ppt/presentation.xml", "<root/>"),
                ("ppt/slides/slide1.xml", "<root><t>Slide text</t></root>"),
            ],
            "Slide text",
        ),
        (
            ".epub",
            [("META-INF/container.xml", "<root/>"), ("chapter.xhtml", "<p>Chapter text</p>")],
            "Chapter text",
        ),
    ],
)
def test_other_office(suffix, entries, expected):
    if suffix == ".odt":
        entries.append(("mimetype", "application/vnd.oasis.opendocument.text"))
    assert expected in text(run(zipped(entries), suffix))


def test_zip_duplicate_paths_and_limits():
    with pytest.warns(UserWarning):
        data = zipped([("a.txt", "one"), ("a.txt", "two"), ("../outside.txt", "bad")])
    result = run(data, ".zip")
    assert [c["data"] for c in result["attachments"][:2]] == [b"one", b"two"]
    assert len({c["part"] for c in result["attachments"]}) == 3
    assert result["attachments"][2]["data"] is None
    assert result["status"] == "partial"
    limited = run(data, ".zip", max_attachments=1)
    assert len(limited["attachments"]) == 1
    assert any("enumerated" in w for w in limited["warnings"])
    bomb = run(zipped([("large.txt", "x" * 100000)]), ".zip", max_file_bytes=50)
    assert bomb["attachments"][0]["data"] is None


def test_zip_symlink():
    entry = zipfile.ZipInfo("link.txt")
    entry.create_system = 3
    entry.external_attr = (stat.S_IFLNK | 0o777) << 16
    result = run(zipped([(entry, "/etc/passwd")]), ".zip")
    assert result["attachments"][0]["data"] is None


def test_tar_members_and_links():
    output = io.BytesIO()
    with tarfile.open(fileobj=output, mode="w") as archive:
        entry = tarfile.TarInfo("folder/a.txt")
        entry.size = 5
        archive.addfile(entry, io.BytesIO(b"hello"))
        link = tarfile.TarInfo("link")
        link.type = tarfile.SYMTYPE
        link.linkname = "/etc/passwd"
        archive.addfile(link)
    result = run(output.getvalue())
    assert result["attachments"][0]["data"] == b"hello"
    assert result["attachments"][1]["data"] is None


@pytest.mark.parametrize(
    "compress,suffix", [(gzip.compress, ".gz"), (bz2.compress, ".bzip"), (lzma.compress, ".xz")]
)
def test_compression_and_expansion_limits(compress, suffix):
    assert run(compress(b"hello"), suffix)["attachments"][0]["data"] == b"hello"
    result = run(compress(b"x" * 100000), suffix, max_file_bytes=100)
    assert result["attachments"][0]["data"] is None


def test_emz_retains_emf_gap():
    result = run(gzip.compress(b"\x01\x00\x00\x00fake EMF"), ".emz")
    assert result["status"] == "partial"
    assert result["attachments"][0]["name"] == "content.emf"


def test_real_7z(tmp_path):
    import libarchive

    with libarchive.file_writer(str(tmp_path / "sample.7z"), "7zip") as archive:
        archive.add_file_from_memory("source.txt", len(b"Seven zip contents"), b"Seven zip contents")
    result = run((tmp_path / "sample.7z").read_bytes(), ".7zip")
    assert result["status"] == "ready", result
    assert result["attachments"][0]["data"] == b"Seven zip contents"


@pytest.mark.parametrize(
    "suffix,data", [(".zip", b"PK\x03\x04broken"), (".rar", b"Rar!\x1a\x07bad"), (".doc", b"broken")]
)
def test_corrupt_formats_fail_visibly(suffix, data):
    assert run(data, suffix)["status"] == "failed"


def test_audio_metadata_is_explicitly_not_transcript():
    output = io.BytesIO()
    with wave.open(output, "wb") as audio:
        audio.setnchannels(1)
        audio.setsampwidth(2)
        audio.setframerate(8000)
        audio.writeframes(b"\x00\x00" * 800)
    result = run(output.getvalue())
    assert result["status"] == "partial"
    assert "sample_rate" in text(result)
    assert any("not transcribed" in w for w in result["warnings"])
    assert result["sections"] == run(output.getvalue())["sections"]


def test_nested_archive_ingest_provenance_and_search(tmp_path):
    root = tmp_path / "sources"
    root.mkdir()
    original = gzip.compress(zipped([("nested/letter", b"Archive searchable quokka")]))
    (root / "bundle.tar.gz").write_bytes(original)
    store = initialize(tmp_path / "state", root, {"ocr": "off"})
    try:
        snapshot = ingest(store)
        api = API(store)
        docs = api.inventory(snapshot)["items"]
        assert len(docs) == 3
        assert api.search(snapshot, "quokka", "literal")["total"] == 1
        assert len([d for d in docs if d["parent"]]) == 2
        assert (root / "bundle.tar.gz").read_bytes() == original
        assert ingest(store) == snapshot
        from evidencekg.reports import verify

        assert verify(store)["ok"]
    finally:
        store.close()


def test_signature_covers_new_parsers():
    result = signature(DEFAULTS)
    assert "openpyxl" in result["versions"]
    assert len(result["implementation_sha256"]) == 64


@pytest.mark.parametrize("suffix", [".xls", ".xlt", ""])
def test_real_legacy_excel(suffix):
    from pathlib import Path

    data = (Path(__file__).parent / "fixtures/formats/legacy.xls").read_bytes()
    result = run(data, suffix)
    assert result["status"] == "partial"
    assert "quokka" in text(result)


@pytest.mark.parametrize("suffix", [".doc", ".dot", ""])
def test_real_legacy_word(suffix):
    from pathlib import Path

    if not shutil.which("antiword"):
        pytest.skip("antiword is not installed")
    data = (Path(__file__).parent / "fixtures/formats/legacy.doc").read_bytes()
    result = run(data, suffix)
    assert result["status"] == "partial", result
    assert "quokka" in text(result)


def test_office_expansion_limit():
    book = Workbook()
    book.active["A1"] = "Some data"
    out = io.BytesIO()
    book.save(out)
    result = run(out.getvalue(), ".xlsx", max_attachment_bytes=100)
    assert result["status"] == "failed"
    assert any("expansion limit" in w for w in result["warnings"])


def test_archive_aggregate_limit():
    result = run(zipped([("one", "1234"), ("two", "5678")]), ".zip", max_attachment_bytes=6)
    assert result["attachments"][0]["data"] == b"1234"
    assert result["attachments"][1]["data"] is None


def test_encrypted_zip(tmp_path):
    if not shutil.which("7z"):
        pytest.skip("7z fixture creator not installed")
    (tmp_path / "text.txt").write_text("Encrypted synthetic text")
    subprocess.run(
        ["7z", "a", "-tzip", "-pfixture", str(tmp_path / "encrypted.zip"), "text.txt"],
        cwd=tmp_path,
        check=True,
        capture_output=True,
    )
    result = run((tmp_path / "encrypted.zip").read_bytes(), ".zip")
    assert result["status"] == "partial"
    assert result["attachments"][0]["data"] is None
    assert "unreadable" in result["attachments"][0]["gap"]


def test_archive_depth_gap_and_immutable_children(tmp_path):
    root = tmp_path / "source"
    root.mkdir()
    (root / "bundle.zip").write_bytes(zipped([("child.zip", zipped([("leaf.txt", "Deep quokka")]))]))
    store = initialize(tmp_path / "state", root, {"max_attachment_depth": 1})
    try:
        snapshot = ingest(store)
        manifest = store.manifest(snapshot)
        assert not manifest["inventory_complete"]
        assert any(d["status"] == "failed" and "depth" in str(d["warnings"]) for d in manifest["documents"])
    finally:
        store.close()


def test_extensionless_email_keeps_attachments():
    from email.message import EmailMessage

    email = EmailMessage()
    email["From"] = "fictional@example.invalid"
    email.set_content("Body")
    email.add_attachment(
        b"Attachment quokka", maintype="application", subtype="octet-stream", filename="data"
    )
    result = run(email.as_bytes())
    assert result["attachments"][0]["data"] == b"Attachment quokka"
