import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
from test_core import drain, empty_result

from evidencekg.config import initialize
from evidencekg.db import dump
from evidencekg.ingest import ingest
from evidencekg.inventory import capture
from evidencekg.reports import verify
from evidencekg.retrieval import API
from evidencekg.review_queue import Queue


def test_1201_sources_full_enumeration_schedule_and_neighbours(tmp_path):
    root = tmp_path / "corpus"
    root.mkdir()
    for i in range(1201):
        (root / f"{i:04}.txt").write_text(
            "Common identifier: ZX-12\n\n" + ("Only" if i == 1200 else f"Document {i}")
        )
    store = initialize(
        tmp_path / "state", root, {"identifiers": [{"namespace": "test", "pattern": r"ZX-12"}]}
    )
    snapshot = ingest(store)
    api = API(store)
    ids, cursor, i = set(), None, 0
    while True:
        page = api.inventory(snapshot, cursor, [73, 300, 1000][i % 3])
        ids.update(d["document_version_id"] for d in page["items"])
        cursor, i = page["next_cursor"], i + 1
        if not cursor:
            break
    assert len(ids) == 1201
    feature = store.one("SELECT * FROM features WHERE kind='identifier'")["id"]
    members, cursor = set(), None
    while True:
        page = api.neighbours(snapshot, feature, cursor=cursor, limit=111)
        members.update(i["id"] for i in page["items"])
        cursor = page["next_cursor"]
        if not cursor:
            break
    assert len(members) == 1201
    queue = Queue(store)
    run = queue.start_review(snapshot, "Find qualifiers", model="fake-v1")["run_id"]
    assert (
        store.db.execute(
            "SELECT count(*) FROM review_tasks WHERE run_id=? AND kind='source'", (run,)
        ).fetchone()[0]
        == 1201
    )
    delivered_text = []
    # All source tasks are actually delivered and validated; graph packet structure tested below.
    while (task := queue.next_review_task(run, "fake-v1")) is not None:
        if task["payload"]["kind"] == "source":
            delivered_text.extend(s["text"] for s in task["payload"]["segments"])
        else:
            assert task["payload"]["segments"]
        queue.submit_review_result(
            run, task["task_id"], task["lease_id"], task["input_sha"], empty_result(task)
        )
    assert len(delivered_text) == 1201 and any(t.endswith("Only") for t in delivered_text)
    status = queue.review_status(run)
    assert status["result_validated"] == status["scheduled"]
    assert not store.rows("SELECT * FROM explicit_links")
    assert verify(store)["ok"]
    store.close()


def test_expiry_retry_and_followup(vault):
    root, store = vault
    (root / "x.txt").write_text("Only")
    snapshot = ingest(store)
    queue = Queue(store)
    run = queue.start_review(snapshot, "Review", max_rounds=1)["run_id"]
    task = queue.next_review_task(run, "fake")
    store.db.execute("UPDATE review_tasks SET lease_until=0 WHERE id=?", (task["task_id"],))
    replacement = queue.next_review_task(run, "fake")
    assert replacement["task_id"] == task["task_id"] and replacement["lease_id"] != task["lease_id"]
    with pytest.raises(ValueError, match="Stale"):
        queue.submit_review_result(
            run, task["task_id"], task["lease_id"], task["input_sha"], empty_result(task)
        )
    result = empty_result(replacement)
    result["unresolved_questions"] = ["Is there a hidden amendment?"]
    queue.submit_review_result(
        run, replacement["task_id"], replacement["lease_id"], replacement["input_sha"], result
    )
    drain(queue, run)
    assert (
        store.db.execute("SELECT count(*) FROM review_tasks WHERE kind='reconsideration'").fetchone()[0] == 1
    )
    assert queue.review_status(run)["followups_open"] == 1
    assert queue.review_status(run)["state"] == "scheduled_work_complete_unresolved"
    assert (
        store.one("SELECT * FROM task_attempts WHERE id=?", (task["lease_id"],))["validation_status"]
        == "expired"
    )


def test_hidden_amendment_joint_packet_and_unrelated(vault):
    root, store = vault
    (root / "rule.txt").write_text("The deadline is 17:00.")
    (root / "amendment.txt").write_text("[[rule.txt]] changes the deadline to 18:00.")
    (root / "unrelated.txt").write_text("An azure heron quietly departed.")
    snapshot = ingest(store)
    queue = Queue(store)
    run = queue.start_review(snapshot, "What is the deadline?")["run_id"]
    packets = [
        json.loads(store.get(t["input_manifest_sha"]))
        for t in store.rows("SELECT * FROM review_tasks WHERE run_id=?", (run,))
    ]
    joint = [
        p
        for p in packets
        if p["descriptor"].get("explicit_link", {}).get("relation_type") == "EXPLICIT_DOCUMENT_REFERENCE"
    ]
    assert joint and all("17:00" in dump(p["segments"]) and "18:00" in dump(p["segments"]) for p in joint)
    docs = {d["path"]: d["document_version_id"] for d in store.manifest(snapshot)["documents"]}
    assert API(store).explain_connection(snapshot, docs["rule.txt"], docs["unrelated.txt"])["items"] == []
    assert len([p for p in packets if p["kind"] == "source"]) == 3


def test_changed_during_capture_and_symlink(vault, monkeypatch):
    root, store = vault
    path = root / "x.txt"
    path.write_text("old")
    original = os.fstat
    calls = 0

    def changed(fd):
        nonlocal calls
        calls += 1
        if calls == 2:
            path.write_text("new changed")
        return original(fd)

    monkeypatch.setattr(os, "fstat", changed)
    with pytest.raises(ValueError, match="changed"):
        capture(root, "x.txt", 1000)
    monkeypatch.setattr(os, "fstat", original)
    (root / "escape").symlink_to("/tmp", target_is_directory=True)
    snapshot = ingest(store)
    manifest = store.manifest(snapshot)
    assert not manifest["inventory_complete"]
    assert any(d["path"] == "escape" and d["status"] == "excluded" for d in manifest["documents"])


def test_split_cross_boundary_literal_and_input_limit(tmp_path):
    root = tmp_path / "corpus"
    root.mkdir()
    original = "A" * 31 + "NOT" + "B" * 70
    (root / "long.txt").write_text(original)
    store = initialize(tmp_path / "state", root, {"segment_chars": 32})
    snapshot = ingest(store)
    assert verify(store)["ok"]
    hit = API(store).search(snapshot, "NOT", "literal")["items"][0]
    assert len(hit["sources"]) == 2
    with pytest.raises(ValueError, match="budget"):
        API(store).read_segments(snapshot, [hit["sources"][0]["segment_id"]], max_output_bytes=1)
    store.close()


def test_parser_network_denied():
    source = "from evidencekg.parsers.isolation import deny_network; deny_network(); import socket; socket.socket()"
    p = subprocess.run([sys.executable, "-c", source], capture_output=True, check=False)
    assert p.returncode != 0 and b"Operation not permitted" in p.stderr


def test_pdf_docx_image_adapters(vault):
    from docx import Document
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.pdfgen.canvas import Canvas

    root, store = vault
    cfg = store.config()
    cfg["tessdata"] = os.environ.get("EVIDENCEKG_TEST_TESSDATA")
    store.db.execute("UPDATE corpora SET configuration_json=?", (dump(cfg),))
    pdf = Canvas(str(root / "native.pdf"))
    pdf.drawString(80, 700, "Native PDF qualifier ONLY")
    pdf.showPage()
    pdf.save()
    doc = Document()
    doc.add_paragraph("Document main text")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header"
    table.cell(1, 0).text = "Evidence"
    doc.save(root / "table.docx")
    image = Image.new("RGB", (1300, 260), "white")
    draw = ImageDraw.Draw(image)
    font = (
        ImageFont.truetype("/usr/share/fonts/TTF/DejaVuSans.ttf", 58)
        if Path("/usr/share/fonts/TTF/DejaVuSans.ttf").exists()
        else ImageFont.load_default(size=58)
    )
    draw.text((40, 70), "Scanned evidence ONLY", fill="black", font=font)
    image.save(root / "scan.png")
    pdf = Canvas(str(root / "mixed.pdf"))
    pdf.drawString(80, 760, "Native header")
    pdf.drawImage(str(root / "scan.png"), 50, 300, 500, 100)
    pdf.showPage()
    pdf.drawImage(str(root / "scan.png"), 50, 300, 500, 100)
    pdf.save()
    snapshot = ingest(store)
    docs = {d["path"]: d for d in store.manifest(snapshot)["documents"]}
    assert all(d["status"] == "partial" for d in docs.values())
    api = API(store)
    assert api.search(snapshot, "Native PDF qualifier ONLY", "literal")["total"] == 1
    assert api.search(snapshot, "Evidence", "literal")["total"] >= 1
    assert api.search(snapshot, "Scanned evidence ONLY", "literal")["total"] >= 1
    for d in docs.values():
        assert d["warnings"]
    assert verify(store)["ok"]


@pytest.mark.asyncio
async def test_sdk_stdio_real_tools(vault):
    from mcp import Client, StdioServerParameters

    root, store = vault
    (root / "source.txt").write_text("Literal qualifier only")
    snapshot = ingest(store)
    params = StdioServerParameters(
        command=sys.executable,
        args=["-m", "evidencekg.cli", "--state", str(store.state), "serve", "--transport", "stdio"],
    )
    async with Client(params, read_timeout_seconds=10) as client:
        tools = await client.list_tools()
        assert "inventory" in {t.name for t in tools.tools}
        result = await client.call_tool("inventory", {"snapshot_id": snapshot, "limit": 1})
        assert not result.is_error
        payload = result.structured_content
        assert payload["total"] == 1
        document = payload["items"][0]["document_version_id"]
        section = await client.call_tool(
            "sections", {"snapshot_id": snapshot, "document_version_id": document}
        )
        ids = [s["id"] for s in section.structured_content["items"]]
        read = await client.call_tool("read_segments", {"snapshot_id": snapshot, "segment_ids": ids})
        assert read.structured_content["items"][0]["text"] == "Literal qualifier only"
        scheduled = await client.call_tool("start_review", {"snapshot_id": snapshot, "question": "Review"})
        run = scheduled.structured_content["run_id"]
        task = (
            await client.call_tool("next_review_task", {"run_id": run, "worker_id": "sdk-test"})
        ).structured_content
        accepted = await client.call_tool(
            "submit_review_result",
            {k: task[k] for k in ("run_id", "task_id", "lease_id", "input_sha")}
            | {"result": empty_result(task)},
        )
        assert accepted.structured_content["accepted"]


@pytest.mark.asyncio
async def test_sdk_enumerates_1201_sources_and_memberships(tmp_path):
    from mcp import Client, StdioServerParameters

    root = tmp_path / "sources"
    root.mkdir()
    for i in range(1201):
        (root / f"{i:04}.txt").write_text("Same evidence")
    store = initialize(tmp_path / "state", root)
    snapshot = ingest(store)
    feature = store.one("SELECT id FROM features WHERE kind='paragraph'")["id"]
    params = StdioServerParameters(
        command=sys.executable, args=["-m", "evidencekg.cli", "--state", str(store.state), "serve"]
    )
    async with Client(params, read_timeout_seconds=30) as client:
        for name, extra, key in [
            ("inventory", {}, "document_version_id"),
            ("neighbours", {"node_id": feature}, "id"),
        ]:
            cursor, seen, pages = None, set(), 0
            while True:
                response = await client.call_tool(
                    name,
                    {
                        "snapshot_id": snapshot,
                        "cursor": cursor,
                        "limit": 701 if pages % 2 == 0 else 223,
                        **extra,
                    },
                )
                assert not response.is_error
                page = response.structured_content
                assert not (seen & {d[key] for d in page["items"]})
                seen.update(d[key] for d in page["items"])
                cursor, pages = page["next_cursor"], pages + 1
                if not cursor:
                    break
            assert len(seen) == 1201
    store.close()
