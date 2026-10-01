"""Deterministic document representations with explicit fidelity limitations."""

from __future__ import annotations

import csv
import io
import json
import posixpath
import re
import subprocess
import zipfile
from xml.etree.ElementTree import ParseError

from .formats import DOCUMENTS, TEXT


def checked_zip(data, cfg):
    archive = zipfile.ZipFile(io.BytesIO(data))
    entries = archive.infolist()
    if len(entries) > 20000 or sum(e.file_size for e in entries) > cfg["max_attachment_bytes"]:
        archive.close()
        raise ValueError("Office archive expansion limit; descendants could not be enumerated")
    if any(e.file_size > cfg["max_file_bytes"] for e in entries):
        archive.close()
        raise ValueError("Office member expansion limit; descendants could not be enumerated")
    return archive


# Retain the existing ceiling, including explicitly stored blank/style-only cells.
_SPREADSHEET_CELL_LIMIT = 1_000_000
# Section/locator JSON is much larger than XML values and is enriched downstream.
_SPREADSHEET_OUTPUT_BYTES = 2_000_000
_SHEET_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


class _SpreadsheetLimit(ValueError):
    pass


class _SpreadsheetReader:
    """Bound actual inflated reads across all selected parts, including metadata."""

    def __init__(self, source, budget, member_limit, member_size):
        self.source = source
        self.budget = budget
        self.remaining = member_limit
        self.unread = member_size

    def read(self, size=-1):
        if self.unread == 0:
            return b""
        allowed = min(self.remaining, self.budget[0])
        if allowed <= 0:
            raise _SpreadsheetLimit("Spreadsheet XML expansion limit reached")
        content = self.source.read(min(16384, allowed, size if size >= 0 else 16384))
        self.remaining -= len(content)
        self.budget[0] -= len(content)
        self.unread -= len(content)
        return content


def _spreadsheet_records(archive, name, tags, budget, member_limit):
    """Yield complete records only; unlink processed nodes to bound tree memory."""
    from defusedxml import ElementTree

    info = archive.getinfo(name)
    with archive.open(info) as source:
        reader = _SpreadsheetReader(source, budget, member_limit, info.file_size)
        stack = []
        record = None
        for event, element in ElementTree.iterparse(
            reader, events=("start", "end"), forbid_dtd=True, forbid_entities=True,
            forbid_external=True,
        ):
            if event == "start":
                stack.append(element)
                if len(stack) > 128:
                    raise _SpreadsheetLimit("Spreadsheet XML nesting limit reached")
                if record is None and element.tag in tags:
                    # Do not interpret extension/metadata lookalikes as cell evidence.
                    if element.tag != _SHEET_NS + "c" or [e.tag for e in stack] == [
                        _SHEET_NS + name for name in ("worksheet", "sheetData", "row", "c")
                    ]:
                        record = element
            else:
                if element is record:
                    yield element
                    record = None
                if record is None:
                    if len(stack) > 1:
                        stack[-2].remove(element)
                    element.clear()
                stack.pop()


def _recover_spreadsheet(data, cfg, result, reason):
    """Recover raw evidence, not displayed Excel values or repaired workbooks."""
    from defusedxml.common import DefusedXmlException

    errors = (ValueError, KeyError, OSError, RuntimeError, zipfile.BadZipFile,
              ParseError, DefusedXmlException)
    warnings = result["warnings"]
    warnings.extend([
        "Spreadsheet raw XML recovery after: " + reason[:500],
        "Raw spreadsheet values only: styles/number formats, dates, percentages, currency, "
        "hidden cells, merged cells and formatting semantics are not interpreted; "
        "style references retained, number_format unknown; original workbook required for visual review",
        "Raw formulas and cached values retained, not calculated; cached values may be stale; "
        "shared/array/data-table formula attributes retained without expansion",
        "Recovery visits smaller worksheets first to preserve coverage before limits; "
        "original worksheet names and XML part coordinates retained",
    ])
    # No use of archive.read(): inflated bytes are charged as they are consumed.
    budget = [cfg["max_attachment_bytes"]]
    output_bytes = 0
    output_limit = min(cfg["max_file_bytes"], _SPREADSHEET_OUTPUT_BYTES)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 20000:
            raise ValueError("Office archive entry expansion limit; recovery not attempted")
        names = [entry.filename for entry in entries]
        if len(names) != len(set(names)):
            raise ValueError("Ambiguous duplicate ZIP member names; spreadsheet recovery refused")

        def records(name, tags):
            return _spreadsheet_records(archive, name, tags, budget, cfg["max_file_bytes"])

        rels = {}
        sheets = []
        strings = []
        package_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
        try:
            for element in records("xl/_rels/workbook.xml.rels", {package_ns + "Relationship"}):
                rel_id = element.get("Id")
                if rel_id in rels:
                    raise ValueError("Duplicate workbook relationship ID")
                target = element.get("Target", "")
                part = posixpath.normpath(posixpath.join("xl", target)) if not target.startswith("/") else target[1:]
                if (element.get("TargetMode") == "External" or not part.startswith("xl/")
                        or "\\" in part or ":" in part):
                    rels[rel_id] = None
                else:
                    rels[rel_id] = (part, element.get("Type", ""))
            for element in records("xl/workbook.xml", {_SHEET_NS + "sheet"}):
                sheets.append((element.get("name"), element.get(_REL_NS + "id")))
        except errors as exc:
            warnings.append("Workbook metadata gap: " + str(exc)[:300])
            result["status"] = "failed"
            return

        shared_parts = [r[0] for r in rels.values() if r and r[1].endswith("/sharedStrings")]
        if len(shared_parts) > 1:
            warnings.append("Ambiguous shared string tables; string references remain unresolved")
        elif shared_parts:
            try:
                for element in records(shared_parts[0], {_SHEET_NS + "si"}):
                    if len(strings) >= _SPREADSHEET_CELL_LIMIT:
                        raise _SpreadsheetLimit("Shared string count limit reached")
                    # Exclude phonetic guides (rPh), which are not cell text.
                    strings.append("".join(t.text or "" for t in
                                           element.findall(_SHEET_NS + "t") +
                                           element.findall(_SHEET_NS + "r/" + _SHEET_NS + "t")))
            except errors as exc:
                warnings.append("Shared string table incomplete; unresolved references are gaps: " + str(exc)[:300])

        def sheet_size(sheet):
            relation = rels.get(sheet[1])
            return archive.getinfo(relation[0]).file_size if relation and relation[0] in names else 0

        for title, rel_id in sorted(sheets, key=sheet_size):
            relation = rels.get(rel_id)
            if not relation or not relation[1].endswith("/worksheet") or not title:
                warnings.append(f"Sheet {title!r}: worksheet relationship missing/unsupported; unread")
                continue
            part = relation[0]
            cells = 0
            recovered = 0
            blanks = 0
            unresolved = 0
            invalid = 0
            last = None
            headers = {}
            try:
                for element in records(part, {_SHEET_NS + "c"}):
                    cells += 1
                    if cells > _SPREADSHEET_CELL_LIMIT:
                        raise _SpreadsheetLimit("Spreadsheet cell limit reached (including explicit blank cells)")
                    coordinate = element.get("r", "")
                    match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]{0,6})", coordinate)
                    if not match:
                        invalid += 1
                        continue
                    column = 0
                    for char in match[1]:
                        column = column * 26 + ord(char) - ord("A") + 1
                    row = int(match[2])
                    if column > 16384 or row > 1048576:
                        invalid += 1
                        continue
                    last = coordinate
                    value = element.find(_SHEET_NS + "v")
                    raw = value.text if value is not None else None
                    formula = element.find(_SHEET_NS + "f")
                    kind = element.get("t", "n")
                    text = raw
                    string_status = None
                    if kind == "s":
                        if raw and raw.isascii() and raw.isdecimal() and len(raw) < 10 and int(raw) < len(strings):
                            text = strings[int(raw)]
                            string_status = "resolved"
                        else:
                            text = "[unresolved shared string]"
                            string_status = "unresolved"
                            unresolved += 1
                    elif kind == "inlineStr":
                        inline = element.find(_SHEET_NS + "is")
                        if inline is not None:
                            text = "".join(t.text or "" for t in
                                           inline.findall(_SHEET_NS + "t") +
                                           inline.findall(_SHEET_NS + "r/" + _SHEET_NS + "t"))
                    if formula is not None:
                        text = "=" + formula.text if formula.text else "[formula reference without expression]"
                    if text is None:
                        blanks += 1
                        continue
                    locator = dict(
                        kind="spreadsheet_cell", sheet=title, part=part, cell=coordinate,
                        row=row, column=column, representation="raw_ooxml",
                        raw_value=raw, raw_type=kind, style_index=element.get("s"),
                        number_format=None, formatting_status="uninterpreted",
                        formula=("=" + formula.text) if formula is not None and formula.text else None,
                        formula_attributes=dict(formula.attrib) if formula is not None else None,
                        cached_value=raw if formula is not None else None,
                        shared_string_status=string_status,
                        header_candidate=headers.get(column) if row > 1 else None,
                        header_status="first_row_candidate_not_confirmed",
                    )
                    section = dict(text=text, locator=locator, modality="native")
                    size = len(json.dumps(section, ensure_ascii=False).encode("utf-8")) + 2
                    if output_bytes + size > output_limit:
                        raise _SpreadsheetLimit(f"Spreadsheet recovered output byte limit reached ({output_limit} bytes)")
                    output_bytes += size
                    result["sections"].append(section)
                    recovered += 1
                    if row == 1 and string_status != "unresolved":
                        headers[column] = text
            except errors as exc:
                warnings.append(
                    f"Sheet {title!r} ({part}): incomplete after {cells - 1 if cells > _SPREADSHEET_CELL_LIMIT else cells} "
                    f"cell records, last coordinate {last!r}, {recovered} recovered cells: {str(exc)[:300]}; "
                    "remaining cells/formatting unread"
                )
            if blanks or unresolved or invalid:
                warnings.append(
                    f"Sheet {title!r}: {blanks} blank/style-only records omitted from text (formatting unreviewed); "
                    f"{unresolved} unresolved shared strings; {invalid} invalid/missing cell coordinates skipped"
                )
        # Failed means no usable evidence, never a misleading empty partial success.
        result["status"] = "partial" if result["sections"] else "failed"


def extract_document(data, suffix, cfg, directory):
    if suffix not in DOCUMENTS | TEXT or suffix in (".docx", ".dotx", ".docm", ".dotm", ".doct"):
        return None
    result = dict(sections=[], warnings=[], attachments=[], artifacts=[], status="ready")

    def add(text, **locator):
        result["sections"].append(dict(text=text, locator=locator, modality="native"))

    def decode():
        encoding = "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        try:
            return data.decode(encoding)
        except UnicodeError:
            result["warnings"].append("Invalid/unknown text encoding; replacement characters retained")
            return data.decode("utf-8", errors="replace")

    if suffix in (".html", ".htm", ".xhtml"):
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(data, "html.parser")
        for node in soup(["script", "style", "template"]):
            node.decompose()
        add(soup.get_text("\n"), kind="html")
        result["warnings"].append("HTML layout, scripts, remote resources and visual content not evaluated")
    elif suffix == ".rtf":
        from striprtf.striprtf import rtf_to_text

        add(rtf_to_text(decode(), errors="replace"), kind="rtf")
        result["warnings"].append(
            "RTF formatting and embedded objects not reviewed; descendants could not be enumerated"
        )
    elif suffix == ".ics":
        # Unfold lines per RFC 5545, preserve properties/parameters/timezones verbatim.
        text = decode().replace("\r\n", "\n").replace("\n ", "").replace("\n\t", "")
        add(text, kind="calendar_properties")
        result["warnings"].append(
            "Calendar properties preserved; recurrence, timezone and alarms not evaluated"
        )
    elif suffix in (".csv", ".tsv"):
        for row, fields in enumerate(
            csv.reader(io.StringIO(decode()), delimiter="\t" if suffix == ".tsv" else ","), 1
        ):
            add(json.dumps(fields, ensure_ascii=False), kind="delimited_row", row=row)
    elif suffix in (".xlsx", ".xltx", ".xlsm", ".xltm", ".xlst"):
        from openpyxl import load_workbook

        try:
            with checked_zip(data, cfg):
                book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
                cached = None
                try:
                    cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
                    for sheet in book:
                        # Ignore declared dimensions: malformed spreadsheets can claim 1M x 16K cells.
                        sheet.reset_dimensions()
                        cached[sheet.title].reset_dimensions()
                        cached_rows = iter(cached[sheet.title].iter_rows())
                        header = []
                        cells = 0
                        for row in sheet.iter_rows():
                            values = next(cached_rows, ())
                            if not header:
                                header = [str(c.value) if c.value is not None else None for c in row]
                            cells += len(row)
                            if cells > 1_000_000:
                                raise ValueError("Spreadsheet cell limit")
                            for cell in row:
                                if cell.value is not None:
                                    add(
                                        str(cell.value),
                                        kind="spreadsheet_cell",
                                        sheet=sheet.title,
                                        cell=cell.coordinate,
                                        row=cell.row,
                                        column=cell.column,
                                        number_format=cell.number_format,
                                        formula=cell.value if cell.data_type == "f" else None,
                                        cached_value=str(values[cell.column - 1].value)
                                        if cell.data_type == "f" and len(values) >= cell.column
                                        and values[cell.column - 1].value is not None else None,
                                        header_candidate=header[cell.column - 1]
                                        if cell.row > 1 and len(header) >= cell.column else None,
                                        header_status="first_row_candidate_not_confirmed",
                                    )
                finally:
                    book.close()
                    if cached is not None:
                        cached.close()
        except (ValueError, TypeError, IndexError, KeyError, ParseError, zipfile.BadZipFile) as exc:
            # Re-extract once from raw XML; never mix interpreted and raw cells.
            result["sections"].clear()
            _recover_spreadsheet(data, cfg, result, str(exc))
        result["warnings"].append(
            "Spreadsheet formulas retained, not calculated; charts, macros and embedded objects unreviewed; descendants could not be enumerated"
        )
    elif suffix in (".xls", ".xlt"):
        import xlrd

        book = xlrd.open_workbook(file_contents=data, on_demand=True)
        try:
            for sheet in book.sheets():
                if sheet.nrows * sheet.ncols > 1_000_000:
                    raise ValueError("Spreadsheet cell limit")
                for row in range(sheet.nrows):
                    for col in range(sheet.ncols):
                        cell = sheet.cell(row, col)
                        if cell.ctype not in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                            value = cell.value
                            if cell.ctype == xlrd.XL_CELL_DATE:
                                value = xlrd.xldate_as_datetime(value, book.datemode).isoformat()
                            add(
                                str(value),
                                kind="spreadsheet_cell",
                                sheet=sheet.name,
                                row=row + 1,
                                column=col + 1,
                            )
        finally:
            book.release_resources()
        result["warnings"].append(
            "Legacy spreadsheet cached values may be stale; macros/charts/objects unreviewed; descendants could not be enumerated"
        )
    elif suffix in (".doc", ".dot"):
        completed = subprocess.run(
            ["antiword", "-m", "UTF-8.txt", str(directory / "input")],
            capture_output=True,
            timeout=min(cfg["parser_timeout"], 60),
        )
        if completed.returncode:
            raise ValueError("Legacy Word extraction failed (encrypted, damaged or unsupported document)")
        add(completed.stdout.decode("utf-8", errors="replace"), kind="legacy_word")
        result["warnings"].append(
            "Legacy Word text only; layout, revisions and embedded objects unreviewed; descendants could not be enumerated"
        )
    elif suffix in (".pptx", ".pptm", ".potx", ".odt", ".ods", ".odp", ".epub"):
        from defusedxml import ElementTree

        with checked_zip(data, cfg) as archive:
            for name in sorted(archive.namelist()):
                if suffix.startswith(".ppt") or suffix == ".potx":
                    selected = name.startswith(
                        ("ppt/slides/slide", "ppt/notesSlides/notesSlide")
                    ) and name.endswith(".xml")
                elif suffix == ".epub":
                    selected = name.lower().endswith((".html", ".xhtml", ".htm"))
                else:
                    selected = name == "content.xml"
                if not selected:
                    continue
                content = archive.read(name)
                if suffix == ".epub":
                    from bs4 import BeautifulSoup

                    soup = BeautifulSoup(content, "html.parser")
                    for node in soup(["script", "style"]):
                        node.decompose()
                    add(soup.get_text("\n"), kind="epub_part", part=name)
                else:
                    root = ElementTree.fromstring(content)
                    # Preserve paragraph and cell boundaries and do not expand repeated ODF cells.
                    tags = {"p", "h"} if suffix in (".odt", ".ods", ".odp") else {"t"}
                    for index, element in enumerate(root.iter()):
                        if element.tag.rsplit("}", 1)[-1] in tags:
                            add("".join(element.itertext()), kind="office_xml", part=name, element=index)
        result["warnings"].append(
            "Office/book text only; layout, repeated cells, formulas, media and embedded objects unreviewed; descendants could not be enumerated"
        )
    else:
        add(decode(), kind="text")
    if result["warnings"] and result["status"] != "failed":
        result["status"] = "partial"
    return result
